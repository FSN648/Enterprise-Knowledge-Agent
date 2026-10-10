from __future__ import annotations

from pathlib import Path
from typing import Any
import hashlib

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from langchain_community.vectorstores import Chroma
from langchain_pinecone import PineconeVectorStore
from pinecone import Pinecone, ServerlessSpec

from .config import get_settings

settings = get_settings()


def _embeddings() -> OpenAIEmbeddings:
    return OpenAIEmbeddings(
        model=settings.openai_embedding_model,
        api_key=settings.openai_api_key,
    )


def _pinecone_store() -> PineconeVectorStore:
    pc = Pinecone(api_key=settings.pinecone_api_key)
    existing = [x["name"] for x in pc.list_indexes()]
    if settings.pinecone_index not in existing:
        pc.create_index(
            name=settings.pinecone_index,
            dimension=1536,
            metric="cosine",
            spec=ServerlessSpec(cloud="aws", region="us-east-1"),
        )
    return PineconeVectorStore.from_existing_index(
        index_name=settings.pinecone_index,
        embedding=_embeddings(),
        namespace=settings.pinecone_namespace,
    )


def _chroma_store() -> Chroma:
    return Chroma(
        collection_name="enterprise_assistant",
        embedding_function=_embeddings(),
        persist_directory=settings.chroma_dir,
    )


def vector_store():
    if settings.vector_backend.lower() == "pinecone":
        if not settings.pinecone_api_key:
            raise RuntimeError("VECTOR_BACKEND=pinecone but PINECONE_API_KEY is missing")
        return _pinecone_store()
    return _chroma_store()


def load_documents(directory: str | None = None) -> list[Document]:
    root = Path(directory or settings.documents_dir)
    root.mkdir(parents=True, exist_ok=True)

    documents: list[Document] = []

    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue

        suffix = path.suffix.lower()

        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8")
            documents.append(Document(
                page_content=text,
                metadata={"source": str(path), "file_name": path.name},
            ))

        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    documents.append(Document(
                        page_content=text,
                        metadata={
                            "source": str(path),
                            "file_name": path.name,
                            "page": page_number,
                        },
                    ))

        elif suffix == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(str(path))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            if text:
                documents.append(Document(
                    page_content=text,
                    metadata={"source": str(path), "file_name": path.name},
                ))

    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=150,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_documents(documents)

    for chunk in chunks:
        source = chunk.metadata.get("source", "unknown")
        digest = hashlib.sha1(
            f"{source}:{chunk.page_content}".encode("utf-8")
        ).hexdigest()
        chunk.metadata["chunk_id"] = digest

    return chunks


def ingest_documents() -> tuple[int, int]:
    docs = load_documents()
    chunks = split_documents(docs)

    if not chunks:
        return 0, 0

    store = vector_store()
    store.add_documents(chunks)
    return len(docs), len(chunks)


def search(query: str, k: int = 5) -> list[dict[str, Any]]:
    store = vector_store()
    results = store.similarity_search_with_score(query, k=k)

    output = []
    for doc, score in results:
        output.append({
            "content": doc.page_content,
            "source": doc.metadata.get("source", "unknown"),
            "page": doc.metadata.get("page"),
            "score": float(score),
            "metadata": doc.metadata,
        })
    return output
